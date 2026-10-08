import os
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging

def export_data_to_excel(data_list, output_directory):
    """
    Exports list of drawing data dicts to a formatted Excel file.
    data_list: list of dicts from extractor.py
    output_directory: folder where Excel file will be saved
    """
    os.makedirs(output_directory, exist_ok=True)
    
    # 1. Create timestamp and filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"Drawing_Data_Extraction_{timestamp}.xlsx"
    file_path = os.path.join(output_directory, file_name)
    
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Drawing Data"
        
        # Enable grid lines explicitly
        ws.views.sheetView[0].showGridLines = True
        
        # 2. Define headers and mappings
        # We will output all extracted fields, mapping them nicely
        headers = [
            "Drawing No.", "Drawing Title", "Revision", "Weight", "Quantity", 
            "Sheet", "Project", "Client", "Date", "Scale", "Drawn By", 
            "Checked By", "Approved By", "Remark", "Status"
        ]
        
        # Write headers
        ws.append(headers)
        
        # Styling definitions
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark navy blue
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Apply header styles
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            
        # Freeze the first row
        ws.freeze_panes = "A2"
        
        # Styling for data cells
        data_font = Font(name="Segoe UI", size=10)
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")
        
        # Border style
        thin_side = Side(border_style="thin", color="D3D3D3")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        # Highlight fills
        error_fill = PatternFill(start_color="FFD2D2", end_color="FFD2D2", fill_type="solid") # Light red
        error_font = Font(name="Segoe UI", size=10, color="9C0006") # Dark red text
        
        status_valid_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Light green
        status_valid_font = Font(name="Segoe UI", size=10, color="006100", bold=True)
        
        status_suspicious_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid") # Light yellow
        status_suspicious_font = Font(name="Segoe UI", size=10, color="9C6500", bold=True)
        
        # 3. Populate rows
        for row_idx, row_data in enumerate(data_list, start=2):
            ws.row_dimensions[row_idx].height = 22
            
            # Combine sheet number and total sheets
            sheet_num = row_data.get("sheet_number", "")
            tot_sheets = row_data.get("total_sheets", "")
            sheet_val = ""
            if sheet_num and tot_sheets:
                sheet_val = f"{sheet_num} of {tot_sheets}"
            elif sheet_num:
                sheet_val = str(sheet_num)
            
            # Map database keys to columns
            col_values = [
                row_data.get("drawing_number", ""),
                row_data.get("drawing_title", ""),
                row_data.get("revision", ""),
                row_data.get("weight", ""),
                row_data.get("quantity", ""),
                sheet_val,
                row_data.get("project_name", ""),
                row_data.get("client_name", ""),
                row_data.get("date", ""),
                row_data.get("scale", ""),
                row_data.get("drawn_by", ""),
                row_data.get("checked_by", ""),
                row_data.get("approved_by", ""),
                row_data.get("remark", ""),
                row_data.get("status", "Valid")
            ]
            
            # Get validation details for color formatting
            validation = row_data.get("_validation", {})
            
            # Helper to map column index back to field key
            col_keys = [
                "drawing_number", "drawing_title", "revision", "weight", "quantity",
                "sheet_number", "project_name", "client_name", "date", "scale", 
                "drawn_by", "checked_by", "approved_by", "remark"
            ]
            
            for col_idx, val in enumerate(col_values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = cell_border
                
                # Check alignment
                if col_idx in [1, 3, 4, 5, 6, 9, 10, 15]: # Numbers, codes, dates, statuses
                    cell.alignment = data_align_center
                else:
                    cell.alignment = data_align_left
                
                # Apply validation highlighting (except for status column)
                if col_idx <= len(col_keys):
                    key = col_keys[col_idx - 1]
                    
                    # For combined Sheet column, check both sheet_number and total_sheets
                    if key == "sheet_number":
                        val_info_1 = validation.get("sheet_number", {"is_valid": True})
                        val_info_2 = validation.get("total_sheets", {"is_valid": True})
                        cell_valid = val_info_1.get("is_valid", True) and val_info_2.get("is_valid", True)
                    else:
                        val_info = validation.get(key, {"is_valid": True})
                        cell_valid = val_info.get("is_valid", True)
                    
                    if not cell_valid or not val:
                        # If required or value fails validation
                        is_required = False
                        if key == "sheet_number":
                            is_required = False # Sheet number is optional in config
                        else:
                            # Search in config FIELD_DEFINITIONS
                            from config import FIELD_DEFINITIONS
                            f_def = next((f for f in FIELD_DEFINITIONS if f["key"] == key), None)
                            if f_def and f_def["validation"].get("required", False):
                                is_required = True
                        
                        if is_required or (val and not cell_valid):
                            cell.fill = error_fill
                            cell.font = error_font
                
                # Format status column
                if col_idx == len(headers):
                    status = val
                    if status == "Valid":
                        cell.fill = status_valid_fill
                        cell.font = status_valid_font
                    elif status == "Suspicious":
                        cell.fill = status_suspicious_fill
                        cell.font = status_suspicious_font
                    else: # Missing
                        cell.fill = error_fill
                        cell.font = error_font
        
        # 4. Auto-adjust column widths based on cell content length
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            # Determine maximum string length in this column
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            
            # Set width with some extra margin padding
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        # Save file
        wb.save(file_path)
        logging.info(f"Excel file successfully generated and saved to {file_path}")
        return file_path
    
    except Exception as e:
        logging.error(f"Failed to export data to Excel: {e}")
        raise e
